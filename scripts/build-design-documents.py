"""Build the two first-check Word documents from the reviewed repository sources.

Requires python-docx and Pillow. Run from any directory; outputs are under docs.
Rendered QA artifacts belong under output/docx-qa and are not deliverables.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/first-check'
QA = ROOT / 'output/docx-qa'
QA.mkdir(parents=True, exist_ok=True)

def clean(s):
    s = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', s)
    return re.sub(r'\*\*([^*]+)\*\*', r'\1', s).replace('`', '')

def doc(title, intro):
    d = Document()
    for st in d.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2), Cm(2)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        st = d.styles[name]
        st.font.name = 'Microsoft YaHei'
        st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.italic = False
    st = d.styles['Normal']
    st.font.size = Pt(10.5)
    st.paragraph_format.line_spacing = 1.18
    st.paragraph_format.space_after = Pt(6)
    for name, size in [('Title', 23), ('Heading 1', 16), ('Heading 2', 12)]:
        d.styles[name].font.size = Pt(size)
    d.core_properties.author = ''
    d.core_properties.title = title
    d.core_properties.subject = 'OOAD 第一次检查 RBAC0'
    d.add_paragraph('统一权限管理系统', 'Subtitle')
    d.add_paragraph(title, 'Title')
    d.add_paragraph('OOAD 第一次检查  |  2026 年 10 月 21 日')
    d.add_paragraph(intro)
    foot = sec.footer.paragraphs[0]
    foot.alignment = 2
    r = foot.add_run()
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    r._r.addnext(field)
    return d

def p(d, text):
    d.add_paragraph(clean(text))

def heading(d, text, page=False):
    h = d.add_heading(text, 1)
    if page: h.paragraph_format.page_break_before = True

def table(d, headers, rows, widths=None):
    t = d.add_table(rows=1, cols=len(headers)); t.autofit = False
    if widths is None: widths = [17/len(headers)]*len(headers)
    for col, w in zip(t.columns, widths): col.width = Cm(w)
    for i, s in enumerate(headers): t.rows[0].cells[i].text = s
    for row in rows:
        for c, s in zip(t.add_row().cells, row): c.text = clean(str(s))
    for j, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if j == 0: trpr.append(OxmlElement('w:tblHeader'))
        for i, c in enumerate(row.cells):
            c.width = Cm(widths[i]); c.vertical_alignment = 1
            pr = c._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:'+side); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9'); borders.append(el)
            pr.append(borders)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EEF3' if j == 0 else 'FFFFFF'); pr.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'80'); el.set(qn('w:type'),'dxa'); margins.append(el)
            pr.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs: run.font.size = Pt(9); run.bold = j == 0
    d.add_paragraph().paragraph_format.space_after = Pt(0)

def md(d, content):
    lines = content.splitlines(); i=0
    while i < len(lines):
        line=lines[i].strip()
        if line.startswith('```'):
            code=[]; i+=1
            while i<len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]); i+=1
            p(d, '\n'.join(code))
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',x) for x in cells): rows.append(cells)
                i+=1
            widths = [3,7,7] if len(rows[0])==3 else None
            table(d, rows[0], rows[1:], widths); continue
        elif line.startswith('#'):
            d.add_heading(re.sub(r'^#+\s*','',line), 2)
        elif line: p(d,line)
        i+=1

source=(OUT/'design.md').read_text(encoding='utf-8')
sections={}
for match in re.finditer(r'^## (\d+)\. ([^\n]+)\n(.*?)(?=^## |\Z)', source, re.M|re.S): sections[int(match[1])]=match[3].strip()

# Database document: relational overview plus a complete physical field dictionary.
db=doc('数据库设计文档', '本文说明 RBAC0 的七张表如何保存账号、角色、权限、授权关系和登录会话，供建库、代码核对和第一次检查讲解使用。表结构依据仓库 MySQL 建表脚本；本次运行验证使用 H2，MySQL 尚未实测。')
heading(db, '1 设计目标与关系')
p(db,'用户通过角色获得权限。用户与角色、角色与权限均为多对多关系，分别使用关联表保存；每项权限属于一个业务系统。登录会话独立保存，使注销后的 JWT 能立即失效。当前不设计角色继承、角色互斥、租户或部门范围。')
table(db,['关系','基数','关联字段'],[
('用户与角色','多对多','sys_user_role.user_id / role_id'),
('角色与权限','多对多','sys_role_permission.role_id / permission_id'),
('系统与权限','一对多','sys_permission.system_id'),
('用户与会话','一对多','sys_login_session.user_id')],[4,2.5,10.5])
im=Image.new('RGB',(1500,620),'white'); draw=ImageDraw.Draw(im)
font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',25)
boxes=[(40,90,390,175,'sys_user 用户'),(575,90,925,175,'sys_user_role 用户角色'),(1110,90,1460,175,'sys_role 角色'),(1110,300,1460,385,'sys_role_permission'),(575,300,925,385,'sys_permission 权限'),(40,300,390,385,'biz_system 业务系统'),(40,510,390,595,'sys_login_session 会话')]
for x1,y1,x2,y2,label in boxes:
    draw.rounded_rectangle((x1,y1,x2,y2),radius=10,fill='#edf2f5',outline='#6b7b86',width=2)
    draw.text(((x1+x2)/2,(y1+y2)/2),label,font=font,fill='black',anchor='mm')
for points, label, pos in [([(390,132),(575,132)],'1 → N',(480,105)), ([(1110,132),(925,132)],'N ← 1',(1015,105)), ([(1285,175),(1285,300)],'1 → N',(1360,240)), ([(925,342),(1110,342)],'1 → N',(1015,315)), ([(390,342),(575,342)],'1 → N',(480,315)), ([(80,175),(20,175),(20,552),(40,552)],'1 → N',(110,450))]:
    draw.line(points,fill='#596975',width=3); draw.text(pos,label,font=font,fill='black',anchor='mm')
im.save(QA/'database-relations.png')
db.add_picture(str(QA/'database-relations.png'),width=Cm(14.5))
p(db,'图 1  表之间的关系。1 → N 表示一条主表记录可以对应多条关联记录；外键位于 N 端。角色权限关联表通过 role_id 和 permission_id 连接两端。')
p(db,'以单表主键识别实体，以用户名和编码防止重复，以外键保证引用存在。把授权关系独立存储，可以在保留用户和角色的同时单独回收授权。管理员标记属于管理平台身份，不属于业务角色。')

sql=(ROOT/'docs/database-mysql.sql').read_text(encoding='utf-8')
tables={m[1]:m[2] for m in re.finditer(r'CREATE TABLE IF NOT EXISTS (\w+) \((.*?)\) ENGINE',sql,re.S)}
meanings={'id':'记录主键','username':'登录用户名','display_name':'用户显示名','password_hash':'BCrypt 密码摘要','administrator':'是否为平台管理员','enabled':'是否启用','user_id':'关联用户','issued_at':'签发时间','expires_at':'到期时间','code':'全局唯一编码','name':'显示名称','description':'角色说明','base_url':'登记的业务系统地址','system_id':'所属业务系统','menu_path':'菜单路径','role_id':'关联角色','permission_id':'关联权限'}
def dictionary(name, title):
    db.add_heading(title+' '+name,2)
    rows=[]
    for line in tables[name].splitlines():
        m=re.match(r'\s*(\w+) (BIGINT|VARCHAR\(\d+\)|BOOLEAN|TIMESTAMP\(6\)) (.*?)(?:,)?$',line)
        if not m: continue
        field,typ,rest=m.groups(); rest=rest.rstrip(',')
        constraints=['非空' if 'NOT NULL' in rest else '可空']
        if field=='id': constraints.append('主键')
        if 'AUTO_INCREMENT' in rest: constraints.append('自增')
        if 'DEFAULT' in rest: constraints.append('默认 '+rest.split('DEFAULT ')[1])
        meaning=meanings[field]
        if name=='sys_login_session' and field=='id': meaning='JWT 的 jti 会话标识 UUID'
        rows.append((field,typ,'，'.join(constraints),meaning))
    table(db,['字段','MySQL 类型','约束与默认值','含义'],rows,[4,3.2,4.6,5.2])

heading(db,'2 账号与登录会话',True)
dictionary('sys_user','2 1 用户表')
p(db,'username 建唯一索引 uk_sys_user_username。未设置密码时 password_hash 为空，账号不能登录。administrator 默认 false，只能由启动初始化创建管理员，普通创建或更新请求不能修改该标记。密码明文不入库。')
dictionary('sys_login_session','2 2 登录会话表')
p(db,'user_id 外键引用 sys_user.id。索引 idx_login_session_user 支持按用户撤销全部会话，idx_login_session_expires 支持按到期时间检索。表中不保存完整 JWT；到期会话会被鉴权拒绝，但自动清理任务尚未实现。')

heading(db,'3 角色与业务系统',True)
dictionary('sys_role','3 1 角色表')
p(db,'code 建唯一索引 uk_sys_role_code，创建后不可通过更新接口修改。角色停用后不参与有效权限计算，但保留授权关联，恢复启用后可重新生效。')
dictionary('biz_system','3 2 业务系统表')
p(db,'code 建唯一索引 uk_biz_system_code。base_url 仅为系统登记信息，不表示已经接入该业务系统。系统停用后，其权限不参与用户的有效权限计算，但用户登录会话仍可有效。')

heading(db,'4 权限与两类授权关系',True)
dictionary('sys_permission','4 1 权限表')
p(db,'code 建全局唯一索引 uk_sys_permission_code；system_id 引用 biz_system.id。权限编码和所属系统创建后固定。权限没有独立 enabled 字段，由角色和系统状态控制有效性。')
dictionary('sys_user_role','4 2 用户角色表')
p(db,'user_id 引用 sys_user.id，role_id 引用 sys_role.id；组合唯一键 uk_sys_user_role(user_id, role_id) 防止同一角色重复分配。')
dictionary('sys_role_permission','4 3 角色权限表')
p(db,'role_id 引用 sys_role.id，permission_id 引用 sys_permission.id；组合唯一键 uk_sys_role_permission(role_id, permission_id) 防止重复授权。')

heading(db,'5 完整性与生命周期',True)
table(db,['操作','事务内的数据变化'],[
('删除用户','撤销全部会话 → 删除用户角色关系 → 删除用户；首次管理员禁止删除。'),
('删除角色','删除角色权限关系和用户角色关系 → 删除角色。'),
('删除业务系统','删除系统下每项权限及对应角色权限关系 → 删除系统登记。'),
('删除权限','删除角色权限关系 → 删除权限。'),
('回收授权','仅删除指定用户角色或角色权限关联，不删除两端实体。'),
('注销或改密','注销只删当前会话；设置密码与撤销该用户全部会话在同一事务中完成。')],[3.5,13.5])
p(db,'建表脚本没有 ON DELETE CASCADE，应用服务负责按引用顺序删除，外键阻止悬空引用。对象删除是物理删除，没有回收站。前端优先提供停用和回收，不提供实体删除按钮。')
p(db,'登录、设置密码、启停和删除用户共用账号行写锁。数据库唯一约束作为并发重复写入的最后防线。H2 已验证事务回滚与登录改密并发流程，MySQL 的锁和排序规则仍需单独验收。')
db.add_heading('5 1 建库与迁移',2)
p(db,'MySQL 参考脚本为 docs/database-mysql.sql，数据库名 unified_permission，字符集 utf8mb4，排序规则 utf8mb4_unicode_ci，引擎 InnoDB。BOOLEAN 在 MySQL 中按其布尔类型别名规则处理。')
p(db,'已有七表数据库补管理员字段使用 docs/migrations/20261021-administrator.sql。CREATE TABLE IF NOT EXISTS 不会为旧表补字段；更早版本新增会话表的迁移为 docs/migrations/20261019-login-sessions.sql。迁移不会把普通账号自动升级成管理员。')
db.add_heading('5 2 演示数据与验证边界',2)
p(db,'demo 配置使用 H2 内存库 first-check 和 create-drop。进程停止后全部演示数据消失；重新执行 scripts/start-demo.ps1 初始化管理员，再执行 python scripts/demo-data.py 恢复演示基线。凭据在运行时输入，不写入文档或仓库。')
p(db,'本阶段已通过 H2 集成测试和本机 HTTP 验收。MySQL 建库与迁移、持久化重启、备份恢复、服务器部署、多实例首次管理员初始化均未验收。数据库设计不包含真实报销业务表，报销接口只返回固定演示数据。')
db.save(OUT/'数据库设计文档.docx')

# Detailed design: remove the database dictionary and focus on functional behavior.
d=doc('详细设计文档','本文说明第一次检查阶段的功能设计、模块职责、访问控制、关键流程和页面交互，并给出运行与验证方法。系统实现 RBAC0，由管理员配置授权，普通用户按实际有效权限访问业务示例。完整表结构另见同目录的数据库设计文档。')
heading(d,'1 需求与总体设计')
md(d,sections[1])
table(d,['层次','职责','实现位置'],[
('界面','中文表单、列表、状态提示和 API 请求','src/main/resources/static/'),
('入口与鉴权','接收请求、参数校验、管理员拦截','Controller 与 ApiAccessConfig'),
('功能服务','管理、授权、登录、权限计算和事务','RbacService 与 JwtTokenService'),
('数据访问','账号、角色、权限、关联及会话持久化','JPA Repository 与实体')],[2.5,7.3,7.2])
p(d,'第一次检查交付界面设计、数据库设计和面向功能的 RBAC0 后端。角色继承、互斥、OA 审批、HR 同步和 Excel 导入不在本阶段范围。报销页面明确展示演示数据，不表示已接入真实报销系统。')

heading(d,'2 管理员与普通用户边界',True)
md(d,sections[3])
p(d,'管理平台的全局管理员标记与业务角色分离。普通用户不能管理账号、设置他人密码、给自己分配角色或通过请求体创建管理员。管理 API 的 GET 查询也受限制，页面隐藏按钮仅用于简化操作，不代替后端鉴权。')

heading(d,'3 面向功能的模块设计',True)
md(d,sections[4])
p(d,'数据层包含用户、角色、业务系统、权限、用户角色、角色权限和会话七张表。设计把账号与授权关系分开维护：取消勾选只回收关联，不删除用户或角色。字段类型、外键和删除顺序见数据库设计文档。')

heading(d,'4 关键流程与错误处理',True)
d.add_heading('4 1 登录和业务访问',2)
for s in ['1. 用户提交用户名和密码，服务端校验账号启用状态与 BCrypt 摘要。登录成功后保存会话并返回 JWT。',
'2. 普通用户请求 GET /api/demo/expenses，后端校验 JWT 签名、有效期、数据库会话及账号状态。任一项不满足即返回 401。',
'3. 查询用户关联的启用角色，合并这些角色的权限，排除已停用系统的权限，并按权限编码去重。',
'4. 并集包含 expense:list 时返回 200 和固定报销演示数据；否则返回 403。管理员访问业务接口也遵守这条规则。']: p(d,s)
d.add_heading('4 2 授权回收与注销',2)
for s in ['1. 管理员取消用户的角色勾选，网页携带管理员 JWT 发出 DELETE 用户角色关联请求。统一拦截器先验证管理员身份。',
'2. 服务在事务中删除指定关联，返回 204；网页重新查询角色与用户有效权限。',
'3. 普通用户使用原 JWT 再次访问时重新计算权限。如果另一个角色仍提供 expense:list，仍返回 200；回收最后一个来源后返回 403。',
'4. 注销只删除当前 JWT 对应的会话，返回 204；旧 JWT 再次访问返回 401。设置密码或停用用户会撤销该用户的全部会话。']: p(d,s)
d.add_heading('4 3 异常与一致性',2)
md(d,sections[6])

heading(d,'5 主要接口约定',True)
p(d,'管理请求在 Authorization 请求头携带 Bearer 管理员令牌；普通业务请求使用该用户自己的令牌。登录响应中的 accessToken 供后续请求使用，令牌不含可被信任的权限快照。完整请求体与 Postman 集合见 docs/rbac-api.md 和 docs/postman/。')
table(d,['方法与路径','输入或作用','访问条件'],[
('POST /api/auth/login','username、password；返回登录令牌','公开登录'),
('POST /api/auth/logout','撤销当前 Bearer 会话，幂等返回 204','可重复注销'),
('GET /api/auth/me','返回当前身份和实时有效权限','有效登录'),
('GET /api/demo/expenses','固定报销示例数据','有效登录与 expense:list'),
('GET / POST /api/users','用户列表／创建普通用户','管理员'),
('PUT /api/users/{id}','更新显示名和启用状态','管理员'),
('PUT /api/users/{id}/password','设置密码并撤销全部旧会话','管理员'),
('GET / PUT / DELETE\n/api/users/{id}/roles[/角色ID]','查询／分配／回收角色；写入必须带角色 ID','管理员'),
('GET /api/users/{id}/permissions','查询指定用户有效权限','管理员'),
('GET / POST /api/roles、systems、permissions','分别查询／创建角色、系统、权限','管理员'),
('PUT / DELETE\n/api/roles、systems、permissions/{id}','实际路径选用其中一种资源；更新／删除对象','管理员'),
('GET / PUT / DELETE\n/api/roles/{id}/permissions[/权限ID]','查询／分配／回收权限；写入必须带权限 ID','管理员')],[7,7.5,2.5])
p(d,'表中 roles、systems、permissions 是三类路径的简写，实际请求分别使用 /api/roles/{id}、/api/systems/{id}、/api/permissions/{id}。用户也有 DELETE /api/users/{id} 接口。用户名、编码及权限所属系统在创建后不可通过更新接口修改。')
p(d,'错误响应包含 error 和中文 message。400 表示参数不合法，401 表示登录失效，403 表示无访问权限，404 表示对象不存在，409 表示重复数据或受保护管理员操作。')

heading(d,'6 前端交互与实际页面',True)
p(d,'前端采用浅色背景、中文标签、左侧导航和右侧表格。管理员看到用户、角色、系统、权限、我的权限与报销入口；普通用户只看到后两项。编辑区在列表下方展开，小屏导航横向排列。')
p(d,'用户页面支持创建、编辑、启停、设置密码、分配和回收角色，以及查询有效权限；角色页面支持配置权限；系统与权限页面维护业务登记信息。加载、空列表、重复输入、必填字段和操作失败均有文字提示。401 清除当前身份并返回登录页，403 保留登录并提示缺少权限。')
for fn,caption in [('03-user-roles.png','图 1  管理员配置用户角色及查看有效权限'),('09-expenses-403.png','图 2  普通用户缺少权限时显示 403')]:
    d.add_picture(str(OUT/'screenshots'/fn),width=Cm(12))
    p(d,caption)
p(d,'以上为 2026 年 10 月 21 日本机 H2 演示环境的真实截图。登录令牌只存在当前标签页内存，刷新整页需重新登录。完整截图位于 docs/first-check/screenshots/。')

heading(d,'7 测试结果与验证边界',True)
p(d,'2026 年 10 月 21 日使用 Windows、Microsoft OpenJDK 21.0.12、Maven 3.9.14 和 H2 执行 mvn clean verify，20 项测试全部通过，失败、错误和跳过均为 0。真实 HTTP 脚本完成 14 项检查。本文整理已有实测记录，不将文档生成视为重新运行测试。')
table(d,['JaCoCo 指标','已覆盖','未覆盖','比例'],[('行','579','95','85.91%'),('分支','77','27','74.04%'),('指令','2284','330','87.38%'),('方法','249','52','82.72%'),('类','47','1','97.92%')],[5,4,4,4])
table(d,['关键场景','实际结果'],[
('管理员配置与普通用户越权','管理员可配置；普通用户直接请求管理接口返回 403，匿名返回 401。'),
('多角色并集与相同权限来源','合并去重；回收一个角色后，另一角色仍提供权限时返回 200。'),
('原 JWT 与权限回收','回收最后一个权限来源后，同一 JWT 请求返回 403。'),
('注销及会话撤销','注销后旧 JWT 返回 401；停用或改密撤销该用户全部会话。'),
('前端输入与数据刷新','必填、重复用户名、短密码与会话失效已有实测；新建或授权后页面重新读取数据。')],[5.5,11.5])
p(d,'JUnit 涵盖 25 类管理路由的匿名与普通用户访问、初始化幂等、请求体提权、管理员保护、事务回滚和并发登录改密。路由的多次断言不重复计为独立测试。未覆盖包括部分 DTO、PermissionCheckResponse 和部分异常分支。')
p(d,'报告：docs/first-check/evidence/jacoco/index.html；测试原始报告：evidence/surefire/；源码摘要：evidence/test-summary.json；HTTP 结果：evidence/http-verification.json。路径均相对 docs/first-check/，第一个除外。')
p(d,'MySQL 建库迁移、持久化重启、服务器部署、HTTPS 和多实例均未验收。生产审计、限流、密码找回、自动过期会话清理未实现。列表无分页，权限计算逐角色查询；当前结论仅支持小规模本地课程演示。权限回收影响后续请求，不强制中断已开始请求。')

heading(d,'8 本地启动与检查演示',True)
p(d,'准备 JDK 21、Maven 3.9 和 Python 3。在项目根目录设置 JAVA_HOME 指向自己的 JDK 21，然后顺序执行以下命令。网页由 Spring Boot 同源提供，无需额外前端构建。')
for s in ['mvn -f third-party/clonefactory/pom.xml clean install','mvn clean verify','.\\scripts\\start-demo.ps1']: p(d,s)
p(d,'首次命令安装仓库自带 CloneFactory。启动脚本提示输入管理员用户名和密码，密码至少 8 个字符且 UTF-8 不超过 72 字节。保持终端运行，浏览器打开 http://127.0.0.1:18080。密码与 JWT 密钥不写入文档。')
p(d,'另开终端执行 python scripts/demo-data.py --verify，按提示输入管理员凭据和普通演示账号临时密码。脚本准备 demo_alice 的两个角色以及无角色的 demo_bob，完成 14 项检查后恢复演示基线。重复执行会重设演示账号密码并撤销其旧会话，不改变管理员密码。')
table(d,['步骤','展示内容与讲解重点'],[
('1 管理员登录','展示用户、角色、系统、权限页面，说明管理身份独立于业务角色。'),
('2 普通用户登录','另开标签页登录 demo_alice，查看两角色的权限并集并请求报销，返回 200。'),
('3 回收一个角色','管理员回收一个角色；普通用户原会话仍为 200，说明重叠权限保留。'),
('4 回收最后角色','原会话再次请求返回 403，说明权限每次读取当前数据库。'),
('5 注销探测','使用页面注销演示验证旧 JWT 返回 401，说明登录与权限是两层检查。'),
('6 展示设计和报告','展示数据库关系、功能调用、JaCoCo 报告及尚未验证的边界。')],[3.5,13.5])
p(d,'停止 demo 进程会清空 H2 内存数据。下次启动重新初始化管理员并运行数据脚本；不要在未知业务库运行演示初始化。运行中的 target JAR 在 Windows 上可能阻止 mvn clean，只停止自己的演示进程。详细命令和逐步讲解见 docs/first-check/local-run.md 与 demo-script.md。')
d.save(OUT/'详细设计文档.docx')
print('Created database and detailed design DOCX documents')


